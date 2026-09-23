"""
Pruebas de la carga del Formulario 7 por Excel (POST /api/formularios/excel):
validacion de filas, precios, techos, aprobado/rechazado y Drive.

No usan los Excel reales de data/ (no estan en el repo): arman planillas
minimas con la misma estructura (hojas "F7 común" / "F7 especial",
encabezados en la fila 10, datos desde la 11) y una base SQLite temporal.

    python -m unittest discover -s tests -v
"""
import base64
import datetime
import io
import json
import os
import sqlite3
import sys
import tempfile
import unittest
import urllib.error
from pathlib import Path
from unittest import mock

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "backend"))
sys.path.insert(0, str(RAIZ / "scripts"))

for _var in ("DRIVE_APPS_SCRIPT_URL", "DRIVE_TOKEN", "REQUIRE_DRIVE_UPLOAD", "FUENTES_HABILITADAS"):
    os.environ.pop(_var, None)
os.environ["FORMULARIO7_DB_PATH"] = str(Path(tempfile.gettempdir()) / "formulario7_test_import.db")

import app as app_module  # noqa: E402
import auth  # noqa: E402
import db  # noqa: E402
import excel_import  # noqa: E402
import integrations  # noqa: E402

ANIO = app_module.ANIO_FISCAL
CATALOGO = {
    # denominacion: (codigo, unidad, precio)
    "ARROZ": ("2.1.1.00435", "KILOGRAMO", 100),
    "YERBA": ("2.1.1.00454", "KILOGRAMO", 250),
    "MASAS": ("2.1.1.00670", "KILOGRAMO", None),
}
ENCABEZADOS = ["Fuente de financiamiento", "Codigo del bien o servicio", "Denominación", "Unidad de medida",
               "Cantidad estimada a adquirir", "Precio estimado por unidad", "Costo estimado Total"]


def comun(fuente, denominacion, cantidad, **pisar):
    """Fila de "F7 común" como la deja Excel: lo que escribe el area (A, C,
    E) mas lo que calcula el VLOOKUP de la planilla (B, D, F, G)."""
    codigo, unidad, precio = CATALOGO.get(denominacion, (0, 0, 0))
    precio = "" if precio is None else precio
    total = cantidad * precio if isinstance(precio, (int, float)) and isinstance(cantidad, (int, float)) else "#VALUE!"
    fila = {"A": fuente, "B": codigo, "C": denominacion, "D": unidad, "E": cantidad, "F": precio, "G": total}
    fila.update(pisar)
    return fila


def especial(fuente, codigo, denominacion, unidad, cantidad, precio, **pisar):
    total = cantidad * precio if isinstance(cantidad, (int, float)) and isinstance(precio, (int, float)) else 0
    fila = {"A": fuente, "B": codigo, "C": denominacion, "D": unidad, "E": cantidad, "F": precio, "G": total}
    fila.update(pisar)
    return fila


def armar_excel(comunes=(), especiales=(), *, sub="1110111000", programa="22 - Atención primaria",
                fecha=datetime.date(2026, 9, 23), hojas=("F7 común", "F7 especial"), encabezados=ENCABEZADOS):
    """Una planilla completa como la deja un area (encabezado con las celdas
    pintadas completas: Subjurisdiccion, Fecha y Programa)."""
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre in hojas:
        ws = libro.create_sheet(nombre)
        ws["A5"], ws["C5"] = "Jurisdiccion:", "Departamento Ejecutivo"
        ws["A6"], ws["C6"] = "Subjurisdiccion: ", sub
        ws["F6"], ws["G6"] = "Fecha", fecha
        ws["A7"] = "Programa o Actividades centrales:" + (f" {programa}" if programa else "")
        for col, texto in enumerate(encabezados, start=1):
            ws.cell(10, col, texto)
        es_comun = "com" in nombre.lower()
        filas = comunes if es_comun else especiales
        for i, fila in enumerate(filas):
            for letra, valor in fila.items():
                ws[f"{letra}{11 + i}"] = valor
        # Filas vacias como las deja Excel: las formulas en 0.
        desde = 11 + len(filas)
        for r in range(desde, desde + 5):
            for col in ((2, 4, 6, 7) if es_comun else (7,)):
                ws.cell(r, col, 0)
        ws.cell(desde + 5, 1, "Subtotal")
        if es_comun:
            ws.cell(desde + 6, 1, "Total general ")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class BaseCarga(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.db_path = Path(self.tmp.name) / "test.db"
        os.environ["FORMULARIO7_DB_PATH"] = str(self.db_path)
        conn = sqlite3.connect(self.db_path)
        conn.executescript((RAIZ / "db" / "schema.sql").read_text(encoding="utf-8"))
        conn.close()
        self._seed()
        cargas = Path(self.tmp.name) / "cargas"
        parche = mock.patch.object(app_module, "CARGAS_DIR", cargas)
        parche.start()
        self.addCleanup(parche.stop)
        self.cargas_dir = cargas
        self.client = app_module.app.test_client()

    def _seed(self):
        with db.conexion() as conn:
            conn.execute("UPDATE fuente_financiamiento SET activa = 1 WHERE id = 131")
            conn.execute("INSERT INTO secretaria (id, nombre, jur, subjurisdiccion) VALUES (1, 'Salud', '04', '1110111000')")
            conn.execute("INSERT INTO secretaria (id, nombre, jur) VALUES (2, 'Otra', '05')")
            self.area = db.crear_usuario(conn, username="salud", password_hash="x", rol="area", secretaria_id=1,
                                         nombre_completo="Area Salud")
            self.otra = db.crear_usuario(conn, username="otra", password_hash="x", rol="area", secretaria_id=2)
            self.admin = db.crear_usuario(conn, username="admin", password_hash="x", rol="admin")
            for sec, cat, fuente, techo in ((1, "22.01.00", 110, 10000), (1, "22.01.00", 131, 5000),
                                            (1, "22.02.00", 110, 20000), (2, "30.01.00", 110, 1000)):
                db.upsert_cuota_categoria(conn, secretaria_id=sec, categoria=cat, fuente=fuente, anio_fiscal=ANIO,
                                          suma_compromiso=None, porcentaje=None, techo=techo)
            for sec, fuente, monto in ((1, 110, 30000), (1, 131, 5000), (2, 110, 1000)):
                db.upsert_secretaria_cuota_total(conn, secretaria_id=sec, fuente=fuente, anio_fiscal=ANIO,
                                                 monto_total=monto)
            for den, (codigo, unidad, precio) in CATALOGO.items():
                db.upsert_catalogo_bien(conn, codigo=codigo, denominacion=den, unidad_texto=unidad, unidad_num=1,
                                        precio=precio, anio_fiscal=ANIO)
            self.tokens = {}
            for nombre, usuario_id in (("area", self.area), ("otra", self.otra), ("admin", self.admin)):
                self.tokens[nombre] = auth.generar_token()
                db.crear_sesion(conn, usuario_id, self.tokens[nombre])

    def subir(self, contenido, categoria="22.01.00", nombre="F7_1110111000_22.01.00.xlsx", quien="area"):
        return self.client.post(
            "/api/formularios/excel",
            data={"categoria": categoria, "archivo": (io.BytesIO(contenido), nombre)},
            headers={"Authorization": f"Bearer {self.tokens[quien]}"},
            content_type="multipart/form-data",
        )

    def consultar(self, sql, *params):
        with db.conexion() as conn:
            return [dict(f) for f in conn.execute(sql, params).fetchall()]

    def celdas_con_error(self, respuesta):
        return {(e["hoja"], e["celda"]) for e in respuesta.get_json()["errores"] if e.get("celda")}


class ExcelValido(BaseCarga):
    def test_aprueba_y_registra_por_fuente(self):
        contenido = armar_excel(
            comunes=[comun(110, "ARROZ", 10), comun(131, "YERBA", 4)],
            especiales=[especial(110, "2.1.1.00670", "MASAS", "KILOGRAMO", 3, 700),
                        especial("110", "9.9.9.99999", "SERVICIO DE CATERING", "CADA UNO", 1, 2500)],
            programa="22 - Atención primaria",
        )
        resp = self.subir(contenido)
        self.assertEqual(resp.status_code, 201, resp.get_json())
        cuerpo = resp.get_json()
        self.assertEqual(cuerpo["estado"], "aprobado")
        self.assertEqual({t["fuente"]: t["total"] for t in cuerpo["totales"]}, {110: 1000 + 2100 + 2500, 131: 1000})

        subs = self.consultar("SELECT * FROM f7_submission ORDER BY fuente")
        self.assertEqual([(s["fuente"], s["estado"]) for s in subs], [(110, "enviado"), (131, "enviado")])
        self.assertEqual(subs[0]["programa"], "22 - Atención primaria")
        self.assertEqual(subs[0]["subjurisdiccion"], "1110111000")
        items = self.consultar("SELECT tipo, denominacion, cantidad, precio_unitario, subtotal, catalogo_id "
                               "FROM f7_item WHERE submission_id = ? ORDER BY orden", subs[0]["id"])
        self.assertEqual([(i["tipo"], i["denominacion"], i["subtotal"]) for i in items],
                         [("comun", "ARROZ", 1000), ("especial", "MASAS", 2100), ("especial", "SERVICIO DE CATERING", 2500)])
        self.assertIsNotNone(items[1]["catalogo_id"])  # MASAS existe en el listado, sin precio
        self.assertIsNone(items[2]["catalogo_id"])

        carga = self.consultar("SELECT * FROM f7_carga_excel")[0]
        self.assertEqual(carga["estado"], "aprobado")
        self.assertIsNone(carga["errores"])
        self.assertTrue(Path(carga["archivo_local"]).exists())
        self.assertEqual(Path(carga["archivo_local"]).read_bytes(), contenido)

    def test_precio_comun_sale_del_catalogo_si_el_archivo_no_trae_valores(self):
        # Un archivo guardado sin calcular formulas: B, D, F y G vacios.
        contenido = armar_excel(comunes=[comun(110, "ARROZ", 7, B=None, D=None, F=None, G=None)])
        resp = self.subir(contenido)
        self.assertEqual(resp.status_code, 201, resp.get_json())
        self.assertEqual(resp.get_json()["total"], 700)

    def test_filas_vacias_de_la_planilla_no_cuentan(self):
        # armar_excel ya deja filas con 0 en las formulas despues de los datos
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]))
        self.assertEqual(resp.status_code, 201, resp.get_json())
        self.assertEqual(resp.get_json()["items"], 1)

    def test_sin_hoja_especial_se_acepta(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)], hojas=("F7 comun",)))
        self.assertEqual(resp.status_code, 201, resp.get_json())


class ExcelConErrores(BaseCarga):
    def assertRechazado(self, resp):
        self.assertEqual(resp.status_code, 422, resp.get_json())
        self.assertEqual(resp.get_json()["estado"], "rechazado")
        self.assertEqual(self.consultar("SELECT * FROM f7_submission"), [])
        self.assertEqual(self.consultar("SELECT estado FROM f7_carga_excel"), [{"estado": "rechazado"}])

    def test_lista_todas_las_celdas_vacias_de_una_vez(self):
        contenido = armar_excel(
            comunes=[comun(110, "ARROZ", 2), comun(110, "YERBA", None, G=0), comun(None, "ARROZ", 1)],
            especiales=[especial(110, None, "SERVICIO", None, 1, None, G=0)],
        )
        resp = self.subir(contenido)
        self.assertRechazado(resp)
        self.assertEqual(self.celdas_con_error(resp), {
            ("F7 común", "E12"), ("F7 común", "A13"),
            ("F7 especial", "B11"), ("F7 especial", "D11"), ("F7 especial", "F11"),
        })
        mensajes = {e["celda"]: e["mensaje"] for e in resp.get_json()["errores"] if e.get("celda")}
        self.assertIn("Cantidad estimada a adquirir", mensajes["E12"])

    def test_precio_comun_distinto_al_de_presupuesto(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 2, F=150, G=300)]))
        self.assertRechazado(resp)
        self.assertEqual(self.celdas_con_error(resp), {("F7 común", "F11")})
        self.assertIn("$ 100,00", resp.get_json()["errores"][0]["mensaje"])

    def test_bien_sin_precio_en_comun_y_bien_con_precio_en_especial(self):
        resp = self.subir(armar_excel(
            comunes=[comun(110, "MASAS", 1)],
            especiales=[especial(110, "2.1.1.00435", "ARROZ", "KILOGRAMO", 1, 90)],
        ))
        self.assertRechazado(resp)
        errores = {(e["hoja"], e["celda"]): e["mensaje"] for e in resp.get_json()["errores"]}
        self.assertIn("F7 especial", errores[("F7 común", "C11")])
        self.assertIn("F7 común", errores[("F7 especial", "C11")])

    def test_denominacion_que_no_esta_o_mal_copiada(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "CAVIAR", 1), comun(110, "arroz ", 1)]))
        self.assertRechazado(resp)
        errores = {e["celda"]: e["mensaje"] for e in resp.get_json()["errores"]}
        self.assertIn("no está en el Listado de bienes", errores["C11"])
        self.assertIn("no está escrito igual", errores["C12"])

    def test_cantidad_con_decimales_cero_o_texto(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1.5), comun(110, "ARROZ", 0),
                                               comun(110, "ARROZ", "muchos")]))
        self.assertRechazado(resp)
        errores = {e["celda"]: e["mensaje"] for e in resp.get_json()["errores"]}
        self.assertIn("entero", errores["E11"])
        self.assertIn("mayor a cero", errores["E12"])
        self.assertIn("número", errores["E13"])

    def test_fuente_invalida_o_sin_techo_en_la_categoria(self):
        resp = self.subir(armar_excel(comunes=[comun(999, "ARROZ", 1), comun(131, "ARROZ", 1)]),
                          categoria="22.02.00", nombre="F7_1110111000_22.02.00.xlsx")
        self.assertRechazado(resp)
        errores = {e["celda"]: e["mensaje"] for e in resp.get_json()["errores"]}
        self.assertIn("no es válida", errores["A11"])
        self.assertIn("no tiene techo presupuestario en la fuente 131", errores["A12"])

    def test_supera_el_techo_de_la_categoria(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 101)]))  # 10.100 > 10.000
        self.assertRechazado(resp)
        techo = [e for e in resp.get_json()["errores"] if e.get("tipo") == "techo"]
        self.assertEqual(len(techo), 1)
        self.assertEqual((techo[0]["nivel"], techo[0]["fuente"], techo[0]["excedente"]), ("categoria", 110, 100))

    def test_supera_el_total_de_la_secretaria(self):
        with db.conexion() as conn:
            conn.execute("UPDATE secretaria_cuota_total SET monto_total = 12000 WHERE secretaria_id = 1 AND fuente = 110")
        ok = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 90)]),  # 9.000 en 22.02.00
                        categoria="22.02.00", nombre="F7_1110111000_22.02.00.xlsx")
        self.assertEqual(ok.status_code, 201, ok.get_json())

        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 50)]))  # 5.000: dentro de su techo de 10.000
        self.assertEqual(resp.status_code, 422)
        techo = [e for e in resp.get_json()["errores"] if e.get("tipo") == "techo"]
        self.assertEqual([(e["nivel"], e["disponible"], e["excedente"]) for e in techo], [("secretaria", 3000, 2000)])

    def test_nombre_de_archivo_de_otra_categoria(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]), nombre="F7_1110111000_22.02.00.xlsx")
        self.assertRechazado(resp)
        self.assertIn("22.02.00", resp.get_json()["errores"][0]["mensaje"])

    def test_archivo_que_no_es_la_planilla(self):
        self.assertRechazado(self.subir(b"esto no es un excel"))

    def test_planilla_sin_las_hojas_del_formulario(self):
        for hojas in (("Hoja1",), ("F7 especial",)):
            resp = self.subir(armar_excel(especiales=[especial(110, "1", "X", "U", 1, 1)], hojas=hojas))
            self.assertEqual(resp.status_code, 422)
            self.assertIn("no tiene la hoja «F7 común»", resp.get_json()["errores"][0]["mensaje"])

    def test_encabezados_corridos(self):
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)], encabezados=["x"] * 7))
        self.assertEqual(resp.status_code, 422)
        self.assertIn("encabezados", resp.get_json()["errores"][0]["mensaje"])

    def test_planilla_vacia(self):
        resp = self.subir(armar_excel())
        self.assertEqual(resp.status_code, 422)
        self.assertIn("ningún bien", resp.get_json()["errores"][0]["mensaje"])


class RequisitosDelInstructivo(BaseCarga):
    """Lo que pide el instructivo de la propia planilla (celda K2)."""

    def errores(self, resp):
        self.assertEqual(resp.status_code, 422, resp.get_json())
        return resp.get_json()["errores"]

    def test_celdas_pintadas_del_encabezado(self):
        errores = self.errores(self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)], sub=None, fecha=None,
                                                      programa=None)))
        por_celda = {e["celda"]: e for e in errores}
        self.assertEqual(set(por_celda), {"C6", "G6", "A7"})
        self.assertEqual(por_celda["C6"]["campo"], "Subjurisdicción")
        self.assertIn("Programa o Actividades centrales", por_celda["A7"]["mensaje"])
        self.assertTrue(all(e["hoja"] == "F7 común" for e in errores))

    def test_subjurisdiccion_de_10_digitos_y_de_la_secretaria(self):
        errores = self.errores(self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)], sub="111-01"),
                                          nombre="F7_111-01_22.01.00.xlsx"))
        self.assertIn("tiene que ser 1110111000 (la de tu Secretaría)", errores[0]["mensaje"])
        self.assertIn("10 dígitos", errores[-1]["mensaje"])
        # Salud tiene cargada 1110111000: otra subjurisdiccion en la planilla no va.
        errores = self.errores(self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)], sub="1110101000"),
                                          nombre="F7_1110101000_22.01.00.xlsx"))
        self.assertEqual([e["celda"] for e in errores], ["C6"])
        self.assertIn("tu Secretaría es 1110111000", errores[0]["mensaje"])

    def test_nombre_del_archivo(self):
        contenido = armar_excel(comunes=[comun(110, "ARROZ", 1)])
        errores = self.errores(self.subir(contenido, nombre="formulario 7 salud.xlsx"))
        self.assertEqual(len(errores), 1)
        self.assertIn("renombralo como F7_1110111000_22.01.00.xlsx", errores[0]["mensaje"])
        errores = self.errores(self.subir(contenido, nombre="F7_1110101000_22.01.00.xlsx"))
        self.assertIn("tiene que ser 1110111000 (la de la celda C6)", errores[0]["mensaje"])
        # La copia que Windows renombra al bajarla dos veces se acepta.
        self.assertEqual(self.subir(contenido, nombre="F7_1110111000_22.01.00 (1).xlsx").status_code, 201)

    def test_especial_copiado_del_listado(self):
        errores = self.errores(self.subir(armar_excel(especiales=[
            especial(110, "2.1.1.00670", "MASAS FINAS", "KILOGRAMO", 1, 10),  # codigo de MASAS, otro nombre
            especial(110, "1.2.3.45678", "MASAS", "KILOGRAMO", 1, 10),        # nombre de MASAS, otro codigo
            especial(110, "2.1.1.00670", "MASAS", "UNIDAD", 1, 10),           # otra unidad
            especial(110, "9.9.9.99999", "SERVICIO DE SONIDO", "CADA UNO", 1, 10),  # fuera del listado: vale
        ])))
        self.assertEqual({e["celda"] for e in errores}, {"C11", "B12", "D13"})


class Reglas(BaseCarga):
    def test_segundo_excel_reemplaza_la_categoria_entera(self):
        primero = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 90), comun(131, "YERBA", 4)]))
        self.assertEqual(primero.status_code, 201, primero.get_json())
        # Mismo techo de 10.000: si sumara la carga anterior (9.000) no entraria.
        segundo = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 95)]))
        self.assertEqual(segundo.status_code, 201, segundo.get_json())
        subs = {s["fuente"]: s for s in self.consultar("SELECT * FROM f7_submission")}
        self.assertEqual((subs[110]["estado"], subs[131]["estado"]), ("enviado", "anulado"))
        total = self.consultar("SELECT SUM(subtotal) AS t FROM f7_item WHERE submission_id = ?", subs[110]["id"])
        self.assertEqual(total[0]["t"], 9500)

    def test_solo_areas_y_solo_sus_categorias(self):
        contenido = armar_excel(comunes=[comun(110, "ARROZ", 1)])
        self.assertEqual(self.subir(contenido, quien="admin").status_code, 403)
        self.assertEqual(self.subir(contenido, quien="otra").status_code, 400)
        self.assertEqual(self.subir(contenido, nombre="f7.xls").status_code, 400)
        sin_token = self.client.post("/api/formularios/excel", data={"categoria": "22.01.00"})
        self.assertEqual(sin_token.status_code, 401)

    def test_mis_categorias_trae_las_dos_fuentes_y_el_ultimo_excel(self):
        # 100.000: supera el techo de la categoria (10.000) y el total de la Secretaria (30.000)
        self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1000)]))
        resp = self.client.get("/api/mis-categorias", headers={"Authorization": f"Bearer {self.tokens['area']}"})
        cuerpo = resp.get_json()
        self.assertEqual(sorted((c["categoria"], c["fuente"]) for c in cuerpo["categorias"]),
                         [("22.01.00", 110), ("22.01.00", 131), ("22.02.00", 110)])
        self.assertEqual(sorted(t["fuente"] for t in cuerpo["totales"]), [110, 131])
        self.assertEqual(cuerpo["ultimas_cargas"]["22.01.00"]["estado"], "rechazado")
        self.assertEqual(cuerpo["ultimas_cargas"]["22.01.00"]["cantidad_errores"], 2)

    def test_panel_del_admin_lista_los_intentos(self):
        self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]))
        self.subir(armar_excel(comunes=[comun(None, "ARROZ", 1)]))
        resp = self.client.get("/api/admin/cargas-excel", headers={"Authorization": f"Bearer {self.tokens['admin']}"})
        self.assertEqual([c["estado"] for c in resp.get_json()["cargas"]], ["rechazado", "aprobado"])


class FuentesOcultas(BaseCarga):
    """La 131 esta oculta (FUENTES_HABILITADAS por defecto = 110). BaseCarga
    la deja activa para las demas pruebas; aca se la oculta como al arrancar."""

    def ocultar_131(self):
        with db.conexion() as conn:
            db.habilitar_fuentes(conn, (110,))

    def get(self, ruta, quien="area"):
        return self.client.get(ruta, headers={"Authorization": f"Bearer {self.tokens[quien]}"}).get_json()

    def test_por_defecto_solo_se_habilita_la_110(self):
        self.assertEqual(app_module.FUENTES_HABILITADAS, (110,))

    def test_la_131_no_aparece_ni_para_el_area_ni_para_el_admin(self):
        self.ocultar_131()
        cuerpo = self.get("/api/mis-categorias")
        self.assertEqual(sorted((c["categoria"], c["fuente"]) for c in cuerpo["categorias"]),
                         [("22.01.00", 110), ("22.02.00", 110)])
        self.assertEqual([t["fuente"] for t in cuerpo["totales"]], [110])
        self.assertEqual([f["id"] for f in self.get("/api/fuentes", quien="admin")["fuentes"]], [110])

    def test_el_excel_no_acepta_filas_de_la_131(self):
        self.ocultar_131()
        resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1), comun(131, "YERBA", 1)]))
        self.assertEqual(resp.status_code, 422)
        errores = {e["celda"]: e["mensaje"] for e in resp.get_json()["errores"]}
        self.assertEqual(list(errores), ["A12"])
        self.assertIn("tiene que ser 110", errores["A12"])

    def test_ocultarla_no_anula_lo_ya_cargado_en_la_131(self):
        primero = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 10), comun(131, "YERBA", 4)]))
        self.assertEqual(primero.status_code, 201, primero.get_json())
        self.ocultar_131()
        segundo = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 20)]))
        self.assertEqual(segundo.status_code, 201, segundo.get_json())
        estados = {s["fuente"]: s["estado"] for s in self.consultar("SELECT fuente, estado FROM f7_submission")}
        self.assertEqual(estados, {110: "enviado", 131: "enviado"})


class Drive(BaseCarga):
    def test_solo_lo_aprobado_va_a_drive(self):
        contenido = armar_excel(comunes=[comun(110, "ARROZ", 1)])
        with mock.patch.object(integrations, "subir_a_drive",
                               return_value={"id": "abc", "webViewLink": "https://drive.test/abc"}) as subir:
            self.assertEqual(self.subir(armar_excel(comunes=[comun(110, "ARROZ", None, G=0)])).status_code, 422)
            subir.assert_not_called()
            resp = self.subir(contenido)
            self.assertEqual(resp.status_code, 201, resp.get_json())
            self.assertTrue(resp.get_json()["drive"])
            subir.assert_called_once_with("F7_1110111000_22.01.00.xlsx", contenido)
        carga = self.consultar("SELECT drive_file_id, drive_link FROM f7_carga_excel WHERE estado = 'aprobado'")[0]
        self.assertEqual(carga, {"drive_file_id": "abc", "drive_link": "https://drive.test/abc"})

    def test_si_drive_falla_no_se_aprueba(self):
        with mock.patch.object(integrations, "subir_a_drive", side_effect=integrations.IntegracionError("sin red")):
            resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]))
        self.assertEqual(resp.status_code, 502)
        self.assertEqual(self.consultar("SELECT * FROM f7_submission"), [])

    def test_drive_obligatorio_sin_configurar(self):
        with mock.patch.dict(os.environ, {"REQUIRE_DRIVE_UPLOAD": "1"}):
            resp = self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]))
        self.assertEqual(resp.status_code, 502)

    def test_pendientes_se_suben_cuando_se_configura_drive(self):
        # Aprobado sin Drive: queda solo en el servidor.
        self.subir(armar_excel(comunes=[comun(110, "ARROZ", 1)]))
        self.subir(armar_excel(comunes=[comun(110, "ARROZ", 2)]))  # reemplaza al anterior
        import subir_pendientes_a_drive
        with mock.patch.dict(os.environ, {"DRIVE_APPS_SCRIPT_URL": "https://script.test/exec", "DRIVE_TOKEN": "x"}),                 mock.patch.object(integrations, "subir_a_drive",
                                  return_value={"id": "d1", "webViewLink": "https://drive.test/d1"}) as subir,                 mock.patch.object(sys, "argv", ["subir_pendientes_a_drive.py", "--anio", str(ANIO)]):
            subir_pendientes_a_drive.main()
        # Solo el ultimo aprobado de la categoria, con el nombre del instructivo.
        subir.assert_called_once()
        self.assertEqual(subir.call_args.args[0], "F7_1110111000_22.01.00.xlsx")
        with db.conexion() as conn:
            self.assertEqual(db.aprobadas_sin_drive(conn, ANIO), [])


class SubidaADrive(unittest.TestCase):
    """subir_a_drive contra un Apps Script falso (sin red)."""

    def setUp(self):
        parche = mock.patch.dict(os.environ, {"DRIVE_APPS_SCRIPT_URL": "https://script.test/exec",
                                              "DRIVE_TOKEN": "clave"})
        parche.start()
        self.addCleanup(parche.stop)
        self.pedidos = []

    def responder(self, texto):
        pedidos = self.pedidos

        class Respuesta(io.BytesIO):
            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

        def urlopen(pedido, timeout=None):
            pedidos.append(pedido)
            if isinstance(texto, Exception):
                raise texto
            return Respuesta(texto.encode("utf-8"))
        return mock.patch("urllib.request.urlopen", urlopen)

    def test_manda_el_archivo_con_la_clave_y_devuelve_el_link(self):
        with self.responder('{"ok": true, "id": "abc", "url": "https://drive.test/abc"}'):
            resultado = integrations.subir_a_drive("F7_1110111000_01.01.00.xlsx", b"datos")
        self.assertEqual(resultado, {"id": "abc", "webViewLink": "https://drive.test/abc"})
        pedido = self.pedidos[0]
        self.assertEqual((pedido.full_url, pedido.get_method()), ("https://script.test/exec", "POST"))
        cuerpo = json.loads(pedido.data)
        self.assertEqual((cuerpo["token"], cuerpo["nombre"]), ("clave", "F7_1110111000_01.01.00.xlsx"))
        self.assertEqual(base64.b64decode(cuerpo["contenido"]), b"datos")

    def test_errores_del_script_o_de_red_no_aprueban(self):
        for respuesta in ('{"ok": false, "error": "clave inválida"}', "<html>Iniciar sesión</html>",
                          urllib.error.URLError("sin red")):
            with self.subTest(respuesta=respuesta), self.responder(respuesta):
                with self.assertRaises(integrations.IntegracionError):
                    integrations.subir_a_drive("F7_1110111000_01.01.00.xlsx", b"datos")

    def test_sin_configurar_queda_en_el_servidor(self):
        with mock.patch.dict(os.environ, {"DRIVE_APPS_SCRIPT_URL": ""}):
            self.assertIsNone(integrations.subir_a_drive("a.xlsx", b"1"))


class Cors(unittest.TestCase):
    def test_acepta_cada_origen_configurado_y_ninguno_mas(self):
        origenes = ("https://nicolasigalbornoz-alt.github.io", "https://formulario7.moron-suministros.workers.dev")
        cliente = app_module.app.test_client()
        with mock.patch.object(app_module, "FRONTEND_ORIGINS", origenes):
            for origen in origenes:
                resp = cliente.options("/api/login", headers={"Origin": origen})
                self.assertEqual(resp.headers["Access-Control-Allow-Origin"], origen)
            resp = cliente.options("/api/login", headers={"Origin": "https://otro-sitio.example"})
            self.assertEqual(resp.headers["Access-Control-Allow-Origin"], origenes[0])
            self.assertEqual(resp.headers["Vary"], "Origin")


class Esquema(unittest.TestCase):
    def test_base_vieja_recibe_las_tablas_nuevas_sin_perder_datos(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = Path(tmp) / "vieja.db"
            conn = sqlite3.connect(ruta)
            conn.executescript("""
                CREATE TABLE secretaria (id INTEGER PRIMARY KEY AUTOINCREMENT, nombre TEXT NOT NULL UNIQUE,
                    jur TEXT, subjurisdiccion TEXT, creado_en TEXT NOT NULL DEFAULT (datetime('now')));
                CREATE TABLE fuente_financiamiento (id INTEGER PRIMARY KEY, nombre TEXT NOT NULL,
                    activa INTEGER NOT NULL DEFAULT 1);
                INSERT INTO fuente_financiamiento VALUES (110, 'Fuente 110', 1), (131, 'Fuente 131', 1);
                INSERT INTO secretaria (nombre) VALUES ('Salud');
            """)
            conn.close()
            with mock.patch.dict(os.environ, {"FORMULARIO7_DB_PATH": str(ruta)}):
                with db.conexion() as conn:
                    db.asegurar_esquema(conn)
                with db.conexion() as conn:
                    tablas = {f[0] for f in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                    self.assertTrue({"sesion", "f7_carga_excel", "f7_submission"} <= tablas)
                    self.assertEqual(conn.execute("SELECT activa FROM fuente_financiamiento WHERE id = 131").fetchone()[0], 1)
                    self.assertEqual(conn.execute("SELECT nombre FROM secretaria").fetchone()[0], "Salud")


class Formato(unittest.TestCase):
    def test_numeros_como_texto_en_formato_argentino(self):
        self.assertEqual(excel_import._numero("1.234,50"), excel_import.Decimal("1234.50"))
        self.assertEqual(excel_import._numero("$ 12.000"), 12000)
        self.assertEqual(excel_import._numero("1.5"), excel_import.Decimal("1.5"))
        with self.assertRaises(ValueError):
            excel_import._numero("#VALUE!")


if __name__ == "__main__":
    unittest.main()
