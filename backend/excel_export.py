"""
Genera el Excel de una carga de Formulario 7 a partir de la plantilla real
(data/formulario7_modelo.xlsx) -- no arma un archivo nuevo desde cero, edita
una COPIA de la plantilla (misma hoja "F7 comun", mismas formulas VLOOKUP
contra "Listado de bienes") para que lo que baja el area sea exactamente el
mismo formato que ya conocen -- el que alimenta el bot de carga a RAFAM
(ver Pagv2/bot-rafam) y el que se manda por mail a dir.presupuesto.

Solo se escriben Fuente, Denominacion y Cantidad por fila (columnas A, C, E
de "F7 comun", filas 11 en adelante) -- Codigo/Unidad/Precio/Costo total
los resuelve la propia formula VLOOKUP de la plantilla al abrir el archivo
en Excel, igual que hace un area cargando a mano. No se toca la hoja
"F7 especial" (queda vacia, tal cual la plantilla): ya no se admiten items
especiales en esta version.

Nota: este modulo NO recalcula las formulas (no depende de LibreOffice) --
el archivo generado tiene los valores en cache vacios hasta que se abre en
Excel de verdad, que recalcula solo al abrir (calculo automatico, el modo
por defecto). Si en algun momento hace falta que un proceso headless lea
los valores ya resueltos sin pasar por Excel, ahi si hace falta agregar un
paso de recalculo -- no esta hecho a proposito para no atar el backend a
una instalacion de LibreOffice en el servidor de produccion.
"""
import io
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "data" / "formulario7_modelo.xlsx"

HOJA = "F7 común"
FILA_INICIO = 11
COL_FUENTE, COL_DENOMINACION, COL_CANTIDAD = 1, 3, 5  # A, C, E


class PlantillaNoEncontradaError(Exception):
    pass


def generar_excel(*, subjurisdiccion, programa, fuente, items):
    """items: lista de dicts con al menos 'denominacion' y 'cantidad'
    (el shape de f7_item). Devuelve un BytesIO listo para enviar."""
    if not PLANTILLA.exists():
        raise PlantillaNoEncontradaError(
            f"Falta la plantilla {PLANTILLA} -- copiarla ahi (ver data/README.md)."
        )

    wb = openpyxl.load_workbook(PLANTILLA)
    ws = wb[HOJA]

    ws["C6"] = subjurisdiccion or ""
    if programa:
        ws["A7"] = f"Programa o Actividades centrales: {programa}"

    fila = FILA_INICIO
    for item in items:
        ws.cell(row=fila, column=COL_FUENTE, value=str(fuente))
        ws.cell(row=fila, column=COL_DENOMINACION, value=item["denominacion"])
        ws.cell(row=fila, column=COL_CANTIDAD, value=item["cantidad"])
        fila += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def nombre_archivo(*, subjurisdiccion, categoria):
    # Mismo criterio que ya usaba el instructivo del Excel para el asunto
    # del mail: F7_Subjurisdiccion_Categoria programatica.
    sub = subjurisdiccion or "SD"
    return f"F7_{sub}_{categoria}.xlsx"
