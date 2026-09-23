"""
Lectura y validacion del Excel del Formulario 7 que sube cada area: la
planilla oficial ("formulario 7 2027.xlsx"), un archivo por Categoria
programatica, con la fuente de financiamiento fila por fila.

Se leen los valores que el archivo ya trae guardados (openpyxl con
data_only): no se abre Excel ni se ejecuta nada del archivo -- ni formulas
ni macros.

Reglas, tal cual el instructivo de la propia planilla:
- "F7 común": bienes con precio de Presupuesto. El area completa Fuente,
  Denominacion y Cantidad; Codigo, Unidad y Precio los completa la planilla
  (VLOOKUP contra "Listado de bienes"). Aca se resuelven contra
  catalogo_bienes (el listado oficial importado a la base): el precio de un
  bien comun lo fija Presupuesto, nunca se toma el del archivo -- y si el
  archivo muestra otro (formula pisada, planilla de otro anio) es un error.
- "F7 especial": bienes sin precio de Presupuesto (o fuera del listado). El
  area completa todo, incluido el precio. Un bien que SI tiene precio de
  Presupuesto no puede ir aca: va en "F7 común".
- Toda fila con algun dato tiene que tener todas sus columnas completas.

No corta en el primer error: junta todos, cada uno con su hoja y celda,
para que el area corrija todo de una vez (y para marcarlos en la copia que
va por mail, ver marcar_errores).
"""
import io
import re
import unicodedata
import warnings
from collections import defaultdict
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from formato import pesos

HOJA_COMUN = "F7 común"
HOJA_ESPECIAL = "F7 especial"
FILA_ENCABEZADO = 10
FILA_INICIO = 11
# La planilla trae 500 filas en "F7 común" y 300 en "F7 especial"; el tope
# es solo contra hojas con formato aplicado hasta la fila 1.048.576.
MAX_FILAS = 5000

LETRAS = ("A", "B", "C", "D", "E", "F", "G")
NOMBRE_COLUMNA = {
    "A": "Fuente de financiamiento",
    "B": "Código del bien o servicio",
    "C": "Denominación",
    "D": "Unidad de medida",
    "E": "Cantidad estimada a adquirir",
    "F": "Precio estimado por unidad",
    "G": "Costo estimado total",
}
# Lo que escribe el area en cada hoja; el resto de las columnas son
# formulas de la planilla.
ENTRADAS = {"comun": ("A", "C", "E"), "especial": ("A", "B", "C", "D", "E", "F")}
# Encabezados de la fila 10 que tienen que estar para confiar en que las
# columnas son las de la planilla oficial (comparados sin tildes).
ENCABEZADOS = {"A": "fuente", "C": "denominacion", "E": "cantidad", "F": "precio"}

CATEGORIA_EN_NOMBRE = re.compile(r"(?<!\d)(\d{2}\.\d{2}\.\d{2})(?!\d)")
# Margen para redondeos de Excel al comparar precio / costo total.
TOLERANCIA = Decimal("0.5")
CENTAVO = Decimal("0.01")


def _normalizar(texto):
    """Minusculas, sin tildes y con los espacios colapsados:
    "F7 Común " y "f7  comun" dan lo mismo."""
    sin_tildes = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")
    return " ".join(sin_tildes.casefold().split())


def _vacio(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _texto(valor):
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _numero(valor):
    """Decimal; None si la celda esta vacia; ValueError si no es un numero.
    Acepta numeros guardados como texto en formato argentino ("1.234,50",
    "$ 12.000") -- Excel los multiplica igual en sus formulas."""
    if _vacio(valor):
        return None
    if isinstance(valor, bool):
        raise ValueError(valor)
    if isinstance(valor, (int, float, Decimal)):
        numero = Decimal(str(valor))
    else:
        texto = str(valor).strip().replace("$", "").replace(" ", "").replace(" ", "")
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
            texto = texto.replace(".", "")
        try:
            numero = Decimal(texto)
        except InvalidOperation:
            raise ValueError(valor) from None
    if not numero.is_finite():
        raise ValueError(valor)
    return numero


def _sin_resolver(valor):
    """Lo que deja una formula de la planilla que no encontro el bien:
    IFERROR(...;0) -> 0, un listado sin precio -> "", o un error de Excel."""
    if isinstance(valor, str):
        return not valor.strip() or valor.strip() == "0" or valor.strip().startswith("#")
    return valor == 0


def _error(hoja, fila, letra, mensaje):
    return {
        "hoja": hoja,
        "fila": fila,
        "columna": letra,
        "celda": f"{letra}{fila}" if hoja and fila and letra else None,
        "campo": NOMBRE_COLUMNA.get(letra),
        "mensaje": mensaje,
    }


def _lista(numeros):
    numeros = [str(n) for n in sorted(numeros)]
    if not numeros:
        return "(ninguna)"
    return numeros[0] if len(numeros) == 1 else ", ".join(numeros[:-1]) + " o " + numeros[-1]


class _Contexto:
    def __init__(self, *, categoria, catalogo, fuentes_categoria, fuentes_activas, anio_fiscal):
        self.categoria = categoria
        self.fuentes_categoria = set(fuentes_categoria)
        self.fuentes_activas = set(fuentes_activas)
        self.anio_fiscal = anio_fiscal
        # En el orden del listado: ante un nombre o codigo repetido gana el
        # primero, igual que el VLOOKUP de la planilla.
        self.por_denominacion = {}
        self.por_codigo = {}
        for bien in catalogo:
            self.por_denominacion.setdefault(_normalizar(bien["denominacion"]), bien)
            self.por_codigo.setdefault(str(bien["codigo"]).strip(), bien)


def _fila_con_datos(tipo, celdas):
    """Una fila "tiene datos" si el area escribio algo en ella. Las formulas
    de la planilla dejan 0 en las filas vacias, eso no cuenta."""
    for letra in LETRAS:
        valor = celdas[letra]
        if _vacio(valor):
            continue
        if letra in ENTRADAS[tipo] or not _sin_resolver(valor):
            return True
    return False


def _validar_fuente(celdas, ctx, error):
    if _vacio(celdas["A"]):
        return None
    try:
        fuente = _numero(celdas["A"])
    except ValueError:
        fuente = None
    fuente = int(fuente) if fuente is not None and fuente == fuente.to_integral_value() else None
    if fuente not in ctx.fuentes_activas:
        error("A", f"La fuente de financiamiento «{_texto(celdas['A'])}» no es válida: "
                   f"tiene que ser {_lista(ctx.fuentes_activas)}.")
        return None
    if fuente not in ctx.fuentes_categoria:
        error("A", f"La categoría {ctx.categoria} no tiene techo presupuestario en la fuente {fuente}: "
                   f"sus bienes solo se pueden cargar con fuente {_lista(ctx.fuentes_categoria)}.")
        return None
    return fuente


def _validar_cantidad(celdas, error):
    """Devuelve la cantidad (Decimal) aunque tenga decimales -- la fila ya
    queda con error, pero sirve igual para estimar el total contra el techo."""
    if _vacio(celdas["E"]):
        return None
    try:
        cantidad = _numero(celdas["E"])
    except ValueError:
        error("E", "La cantidad tiene que ser un número.")
        return None
    if cantidad <= 0:
        error("E", "La cantidad tiene que ser mayor a cero.")
        return None
    if cantidad != cantidad.to_integral_value():
        error("E", "La cantidad tiene que ser un número entero (sin decimales).")
    return cantidad


def _validar_comun(celdas, denominacion, ctx, error):
    """Bien y precio de Presupuesto de una fila de "F7 común". Devuelve
    (bien del catalogo o None, precio o None)."""
    if denominacion is None:
        return None, None
    bien = ctx.por_denominacion.get(_normalizar(denominacion))
    if bien is None:
        error("C", f"«{denominacion}» no está en el Listado de bienes. Copiá el nombre exacto de la hoja "
                   f"«Listado de bienes»; si el bien no figura, cargalo en la hoja «{HOJA_ESPECIAL}» con su precio.")
        return None, None

    # Lo que muestra el propio archivo en las columnas que completa la
    # planilla (None = se guardo sin calcular: no hay nada que comparar,
    # manda el catalogo).
    codigo_archivo, precio_archivo = celdas["B"], celdas["F"]
    coincide_en_archivo = True
    if codigo_archivo is not None and _sin_resolver(codigo_archivo):
        coincide_en_archivo = False
        error("C", f"La planilla no completó sola el código y la unidad de «{denominacion}»: no está escrito "
                   f"igual que en el Listado de bienes («{bien['denominacion']}»). Copialo y pegalo desde esa hoja.")
    elif codigo_archivo is not None and _texto(codigo_archivo) != str(bien["codigo"]).strip():
        coincide_en_archivo = False
        error("B", f"El código que muestra la planilla ({_texto(codigo_archivo)}) no es el del Listado de bienes "
                   f"vigente para «{bien['denominacion']}» ({bien['codigo']}): usá la planilla oficial del "
                   f"Formulario 7 {ctx.anio_fiscal}.")

    if bien["precio"] is None:
        error("C", f"«{bien['denominacion']}» no tiene precio de Presupuesto: va en la hoja «{HOJA_ESPECIAL}», "
                   f"con el precio que estimen ustedes.")
        return bien, None

    precio = Decimal(str(bien["precio"]))
    if coincide_en_archivo and precio_archivo is not None:
        if _sin_resolver(precio_archivo):
            error("F", f"La planilla no muestra el precio de Presupuesto de «{bien['denominacion']}» "
                       f"({pesos(precio)}): usá la planilla oficial vigente y no borres la fórmula de la columna "
                       f"Precio.")
        else:
            try:
                en_archivo = _numero(precio_archivo)
            except ValueError:
                en_archivo = None
            if en_archivo is None or abs(en_archivo - precio) > TOLERANCIA:
                dice = _texto(precio_archivo) if en_archivo is None else pesos(en_archivo)
                error("F", f"El precio de «{bien['denominacion']}» lo fija Presupuesto: {pesos(precio)}. "
                           f"La planilla dice {dice} (¿se modificó la columna Precio o es una planilla de otro año?).")
    return bien, precio


def _validar_especial(celdas, denominacion, cantidad, ctx, error):
    """Precio que pone el area en una fila de "F7 especial". Devuelve
    (bien del catalogo -- sin precio -- o None, precio o None)."""
    precio = None
    if not _vacio(celdas["F"]):
        try:
            precio = _numero(celdas["F"])
        except ValueError:
            error("F", "El precio tiene que ser un número.")
        else:
            if precio <= 0:
                error("F", "El precio tiene que ser mayor a cero.")
                precio = None

    # Un bien con precio de Presupuesto no puede ir aca: el precio de los
    # comunes lo fija Presupuesto, no el area.
    codigo = None if _vacio(celdas["B"]) else _texto(celdas["B"])
    bien = ctx.por_codigo.get(codigo) if codigo else None
    if bien is None and denominacion:
        bien = ctx.por_denominacion.get(_normalizar(denominacion))
    if bien is not None and bien["precio"] is not None:
        error("C", f"«{bien['denominacion']}» tiene precio de Presupuesto ({pesos(bien['precio'])}): va en la hoja "
                   f"«{HOJA_COMUN}», no en «{HOJA_ESPECIAL}».")
        bien = None

    # La columna G es la formula cantidad x precio de la planilla: si
    # muestra otra cosa, el total del propio archivo no es el real.
    costo = celdas["G"]
    if cantidad is not None and precio is not None and not _vacio(costo):
        if isinstance(costo, str) and costo.strip().startswith("#"):
            error("G", f"La planilla no pudo calcular el costo total ({costo.strip()}): revisá que la cantidad y "
                       f"el precio sean números.")
        else:
            try:
                costo = _numero(costo)
            except ValueError:
                costo = None
            if costo is None or abs(costo - cantidad * precio) > TOLERANCIA:
                error("G", f"El costo total tiene que ser cantidad × precio ({pesos(cantidad * precio)}): no "
                           f"modifiques la fórmula de esa columna.")
    return bien, precio


def _validar_fila(tipo, hoja, fila, celdas, ctx):
    """Devuelve (errores, item valido o None, (fuente, subtotal) o None).
    El subtotal se calcula aunque la fila tenga otros errores, siempre que
    haya fuente, cantidad y precio: sirve para avisar ya del techo."""
    errores = []

    def error(letra, mensaje):
        errores.append(_error(hoja, fila, letra, mensaje))

    for letra in ENTRADAS[tipo]:
        if _vacio(celdas[letra]):
            error(letra, f"Falta completar «{NOMBRE_COLUMNA[letra]}».")

    fuente = _validar_fuente(celdas, ctx, error)
    cantidad = _validar_cantidad(celdas, error)
    denominacion = None if _vacio(celdas["C"]) else _texto(celdas["C"])
    if tipo == "comun":
        bien, precio = _validar_comun(celdas, denominacion, ctx, error)
        codigo = bien["codigo"] if bien else None
        unidad = bien["unidad_texto"] if bien else None
        denominacion = bien["denominacion"] if bien else denominacion
    else:
        bien, precio = _validar_especial(celdas, denominacion, cantidad, ctx, error)
        codigo = None if _vacio(celdas["B"]) else _texto(celdas["B"])
        unidad = None if _vacio(celdas["D"]) else _texto(celdas["D"])

    subtotal = None
    if fuente is not None and cantidad is not None and precio is not None:
        subtotal = (cantidad * precio).quantize(CENTAVO)

    item = None
    if not errores:
        item = {
            "tipo": tipo,
            "fuente": fuente,
            "catalogo_id": bien["id"] if bien else None,
            "codigo": codigo,
            "denominacion": denominacion,
            "unidad_medida": unidad,
            "cantidad": int(cantidad),
            "precio_unitario": float(precio),
            "subtotal": float(subtotal),
            "hoja": hoja,
            "fila": fila,
        }
    return errores, item, (fuente, subtotal) if subtotal is not None else None


def _leer_hoja(ws, tipo, ctx, resultado):
    maximo = min(ws.max_row or (FILA_INICIO + MAX_FILAS), FILA_INICIO + MAX_FILAS)
    filas = ws.iter_rows(min_row=FILA_INICIO, max_row=maximo, max_col=len(LETRAS), values_only=True)
    for fila, valores in enumerate(filas, start=FILA_INICIO):
        celdas = dict(zip(LETRAS, tuple(valores) + (None,) * (len(LETRAS) - len(valores))))
        etiqueta = _normalizar(celdas["A"]) if isinstance(celdas["A"], str) else ""
        if etiqueta.startswith("subtotal") or etiqueta.startswith("total"):
            break  # fin de la tabla de la planilla
        if not _fila_con_datos(tipo, celdas):
            continue
        resultado["filas_con_datos"] += 1

        errores, item, parcial = _validar_fila(tipo, ws.title, fila, celdas, ctx)
        resultado["errores"].extend(errores)
        if item is not None:
            resultado["items"].append(item)
        if parcial is not None:
            fuente, subtotal = parcial
            resultado["totales_por_fuente"][fuente] += subtotal


def _encabezado_valido(fila_encabezado):
    celdas = dict(zip(LETRAS, tuple(fila_encabezado) + (None,) * len(LETRAS)))
    return all(esperado in _normalizar(celdas[letra] or "") for letra, esperado in ENCABEZADOS.items())


def _datos_de_encabezado(filas):
    """Subjurisdiccion (C6) y "Programa o Actividades centrales" (A7, a
    continuacion del rotulo o en alguna celda de esa fila)."""
    subjurisdiccion = programa = None
    if len(filas) >= 6 and len(filas[5]) >= 3 and not _vacio(filas[5][2]):
        subjurisdiccion = _texto(filas[5][2])
    if len(filas) >= 7 and filas[6]:
        rotulo = str(filas[6][0] or "")
        resto = re.sub(r"^\s*programa o actividades centrales\s*:?\s*", "", rotulo, flags=re.IGNORECASE).strip()
        if resto:
            programa = _texto(resto)
        else:
            programa = next((_texto(v) for v in filas[6][1:] if not _vacio(v)), None)
    return subjurisdiccion, programa


def leer_formulario(contenido, *, nombre_archivo, categoria, catalogo, fuentes_categoria, fuentes_activas,
                    anio_fiscal):
    """Lee y valida el Excel subido para `categoria`.

    catalogo: filas de catalogo_bienes (db.listar_catalogo).
    fuentes_categoria: fuentes en las que la Categoria tiene techo.
    fuentes_activas: fuentes habilitadas en general (110, 131).

    Devuelve un dict con `items` (filas validas, cada una con su fuente),
    `errores` (todos, con hoja/celda), `totales_por_fuente` ({fuente:
    Decimal}, para validar techos), `subjurisdiccion` y `programa`."""
    resultado = {
        "items": [],
        "errores": [],
        "totales_por_fuente": defaultdict(Decimal),
        "subjurisdiccion": None,
        "programa": None,
        "filas_con_datos": 0,
    }
    ctx = _Contexto(categoria=categoria, catalogo=catalogo, fuentes_categoria=fuentes_categoria,
                    fuentes_activas=fuentes_activas, anio_fiscal=anio_fiscal)

    # El archivo es uno por Categoria y se nombra con ella
    # (F7_Subjurisdiccion_Categoria): si el nombre dice otra, es otro Excel.
    codigos = set(CATEGORIA_EN_NOMBRE.findall(nombre_archivo or ""))
    if codigos and categoria not in codigos:
        resultado["errores"].append(_error(None, None, None,
            f"El nombre del archivo («{nombre_archivo}») es de la categoría {_lista(codigos)}, pero elegiste la "
            f"{categoria}. Revisá que estés subiendo el Excel de esta categoría."))

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # validaciones de datos / encabezados que openpyxl no soporta
            libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        resultado["errores"].append(_error(None, None, None,
            "No se pudo leer el archivo como Excel (.xlsx). Subí la planilla oficial del Formulario 7 "
            "guardada como .xlsx."))
        return resultado

    try:
        hojas = {_normalizar(nombre): nombre for nombre in libro.sheetnames}
        encontradas = [(hojas.get(_normalizar(oficial)), tipo)
                       for oficial, tipo in ((HOJA_COMUN, "comun"), (HOJA_ESPECIAL, "especial"))]
        if not any(nombre for nombre, _ in encontradas):
            resultado["errores"].append(_error(None, None, None,
                f"El archivo no es la planilla del Formulario 7: no tiene las hojas «{HOJA_COMUN}» ni "
                f"«{HOJA_ESPECIAL}»."))
            return resultado

        for nombre, tipo in encontradas:
            if nombre is None:
                continue
            ws = libro[nombre]
            encabezado = list(ws.iter_rows(min_row=1, max_row=FILA_ENCABEZADO, max_col=len(LETRAS), values_only=True))
            if len(encabezado) < FILA_ENCABEZADO or not _encabezado_valido(encabezado[FILA_ENCABEZADO - 1]):
                resultado["errores"].append(_error(nombre, FILA_ENCABEZADO, None,
                    f"La hoja «{nombre}» no tiene los encabezados del Formulario 7 en la fila {FILA_ENCABEZADO} "
                    f"(Fuente, Código, Denominación, Unidad, Cantidad, Precio, Costo total). Usá la planilla "
                    f"oficial sin agregar ni borrar filas o columnas arriba de la tabla."))
                continue
            if resultado["subjurisdiccion"] is None and resultado["programa"] is None:
                resultado["subjurisdiccion"], resultado["programa"] = _datos_de_encabezado(encabezado)
            _leer_hoja(ws, tipo, ctx, resultado)
    finally:
        libro.close()

    if resultado["filas_con_datos"] == 0 and not resultado["errores"]:
        resultado["errores"].append(_error(None, None, None,
            f"El Excel no tiene ningún bien cargado (las hojas «{HOJA_COMUN}» y «{HOJA_ESPECIAL}» están vacías)."))
    return resultado


# ================= Copia con los errores marcados (para el mail) =================

ROJO = PatternFill(fill_type="solid", fgColor="FFFFC7CE")


def _celda_total_general(libro):
    """Donde anotar los errores que no son de una celda (techos, nombre del
    archivo): el "Total general" de F7 común, o su "Subtotal"."""
    for nombre in libro.sheetnames:
        if _normalizar(nombre) != _normalizar(HOJA_COMUN):
            continue
        ws = libro[nombre]
        subtotal = None
        ultima = min(ws.max_row, FILA_INICIO + MAX_FILAS + 5)
        for (celda,) in ws.iter_rows(min_row=FILA_INICIO, max_row=ultima, max_col=1):
            etiqueta = _normalizar(celda.value) if isinstance(celda.value, str) else ""
            if etiqueta.startswith("total general"):
                return nombre, f"G{celda.row}"
            if etiqueta.startswith("subtotal") and subtotal is None:
                subtotal = (nombre, f"G{celda.row}")
        return subtotal or (nombre, "A1")
    return (libro.sheetnames[0], "A1") if libro.sheetnames else None


def marcar_errores(contenido, errores):
    """Copia del Excel subido con cada celda con error pintada y con un
    comentario que dice que esta mal. Va adjunta al mail de rechazo -- a
    la persona que lo subio no se le devuelve ningun archivo, ve la misma
    lista en la pagina. Devuelve bytes, o None si no se pudo armar (el
    mail sale igual, con los errores en el cuerpo)."""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            libro = openpyxl.load_workbook(io.BytesIO(contenido))
    except Exception:
        return None

    por_celda = defaultdict(list)
    generales = []
    for error in errores:
        if error.get("hoja") in libro.sheetnames and error.get("celda"):
            por_celda[(error["hoja"], error["celda"])].append(error["mensaje"])
        else:
            generales.append(error["mensaje"])
    if generales:
        destino = _celda_total_general(libro)
        if destino:
            por_celda[destino].extend(generales)

    for (hoja, celda), mensajes in por_celda.items():
        try:
            destino = libro[hoja][celda]
            destino.fill = ROJO
            comentario = Comment("\n".join(f"• {m}" for m in mensajes), "Formulario 7")
            comentario.width = 380
            comentario.height = min(600, 50 + 50 * len(mensajes))
            destino.comment = comentario
        except (AttributeError, ValueError, KeyError):
            continue  # celda combinada o fuera de la hoja: queda solo en la lista del mail

    salida = io.BytesIO()
    try:
        libro.save(salida)
    except Exception:
        return None
    return salida.getvalue()
