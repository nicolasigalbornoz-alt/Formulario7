"""
Importa el catalogo de bienes ("formulario 7 2027.xlsx", hoja "Listado de
bienes") a catalogo_bienes: Denominacion, Codigo, Unidad de medida y
Precios 2026 (el precio a usar como estimacion para el Formulario 7 -- ver
db/README.md, seccion "Por que anio_fiscal 2027 y no 2026").

Bienes sin precio asignado (precio = NULL) se importan igual, no se
descartan: el formulario los necesita para autocompletar Codigo/Unidad al
cargarlos como "especial" (instructivo, punto 2), solo que no pueden
cargarse como "comun".

Codigo duplicado conocido (2.2.2.04072 / GORRO, dos filas): se queda con la
primera coincidencia, igual que el VLOOKUP del Excel original.

Idempotente, mismo criterio que build_cuota_data.py: upsert por
(codigo, anio_fiscal), lo que no reaparece en una reimportacion se marca
vigente=0 en vez de borrarse.

Uso:
    python scripts/build_catalogo_data.py --anio 2027
"""
import argparse
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "backend"))
import db  # noqa: E402

ARCHIVO_DEFECTO = RAIZ / "data" / "formulario 7 2027.xlsx"
HOJA = "Listado de bienes"
COL_DENOMINACION, COL_CODIGO, COL_UNIDAD_TEXTO, COL_UNIDAD_NUM, COL_PRECIO = 1, 2, 3, 4, 7  # A,B,C,D,G
FILA_INICIO = 2


def leer_bienes(ws):
    filas = []
    r = FILA_INICIO
    while True:
        denominacion = ws.cell(row=r, column=COL_DENOMINACION).value
        if denominacion is None or str(denominacion).strip() == "":
            break
        codigo = ws.cell(row=r, column=COL_CODIGO).value
        precio = ws.cell(row=r, column=COL_PRECIO).value
        filas.append({
            "denominacion": str(denominacion).strip(),
            "codigo": str(codigo).strip() if codigo is not None else None,
            "unidad_texto": str(ws.cell(row=r, column=COL_UNIDAD_TEXTO).value or "").strip(),
            "unidad_num": ws.cell(row=r, column=COL_UNIDAD_NUM).value,
            "precio": round(float(precio), 2) if isinstance(precio, (int, float)) else None,
        })
        r += 1
    return filas


def importar(anio_fiscal, archivo, conn):
    wb = openpyxl.load_workbook(archivo, data_only=True)
    ws = wb[HOJA]
    bienes = leer_bienes(ws)

    db.marcar_catalogo_no_vigente(conn, anio_fiscal)

    vistos = set()
    duplicados = []
    sin_precio = 0
    for b in bienes:
        if not b["codigo"]:
            continue  # fila sin codigo, no se puede cargar como "comun" -- se ignora
        if b["codigo"] in vistos:
            duplicados.append(b["codigo"])
            continue  # primera coincidencia gana, igual que VLOOKUP
        vistos.add(b["codigo"])
        if b["precio"] is None:
            sin_precio += 1
        db.upsert_catalogo_bien(
            conn, codigo=b["codigo"], denominacion=b["denominacion"], unidad_texto=b["unidad_texto"],
            unidad_num=b["unidad_num"], precio=b["precio"], anio_fiscal=anio_fiscal,
        )

    return {"total": len(vistos), "sin_precio": sin_precio, "duplicados": duplicados}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--anio", type=int, required=True, help="Anio fiscal del formulario (ej. 2027)")
    parser.add_argument("--archivo", type=Path, default=ARCHIVO_DEFECTO)
    args = parser.parse_args()

    if not args.archivo.exists():
        print(f"No se encontro {args.archivo}. Copiar ahi el Excel del formulario (ver data/README.md).")
        raise SystemExit(1)

    with db.conexion() as conn:
        resumen = importar(args.anio, args.archivo, conn)

    print(f"Anio {args.anio}: {resumen['total']} bienes importados "
          f"({resumen['sin_precio']} sin precio -- solo cargables como 'especial').")
    if resumen["duplicados"]:
        print(f"  Aviso: {len(resumen['duplicados'])} codigo(s) duplicado(s), se uso la primera fila: "
              f"{sorted(set(resumen['duplicados']))}")


if __name__ == "__main__":
    main()
